import json
import os
from pathlib import Path
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from product_qa.checks import check_product
from product_qa.doctor import tool_path, java_path

BASE = Path(__file__).resolve().parents[1] / "products"


def build_prompt_pack() -> Path:
    folder = BASE / "prompts-fashion-brand"
    folder.mkdir(parents=True, exist_ok=True)

    metadata = {
        "title": "Paket Prompt Pemasaran dan Peluncuran Clothing Brand",
        "author": "Digital Product Factory",
        "version": "1.0.0",
        "license": "CC0-1.0",
    }
    (folder / "product.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    prompts = [
        {
            "id": "deskripsi-katalog-tshirt",
            "template": "Tuliskan deskripsi produk katalog untuk produk {{nama_produk}} dari brand {{nama_brand}}. Produk ini menggunakan bahan {{jenis_bahan}} dengan cutting {{tipe_cutting}}. Sorot kenyamanan pemakaian sehari-hari, detail jahitan rantai, dan panduan padu padan (styling) untuk gaya kasual.",
            "variables": ["nama_produk", "nama_brand", "jenis_bahan", "tipe_cutting"],
            "example": {
                "input": {
                    "nama_produk": "Heavyweight Boxy T-Shirt",
                    "nama_brand": "Arka Apparel",
                    "jenis_bahan": "Cotton Combed 16s (230 gsm)",
                    "tipe_cutting": "Boxy Oversized Fit",
                },
                "output": "Hadirkan siluet tegas dan kenyamanan maksimal bersama Arka Heavyweight Boxy T-Shirt. Dibuat dari Cotton Combed 16s berbobot 230 gsm, kaos ini menjaga bentuknya tetap kokoh tanpa terasa kaku di iklim tropis. Dilengkapi jahitan kerah rib anti-molor dan jahitan rantai pundak berkualitas garment. Padukan bersama celana kargo atau denim lurus untuk siluet streetwear yang rapi.",
            },
        },
        {
            "id": "hook-konten-tiktok-fashion",
            "template": "Buat 5 variasi kalimat hook video TikTok berdurasi 3 detik pertama untuk mempromosikan {{nama_produk}} bertema {{tema_konten}}. Target audiens adalah {{target_audiens}} yang mengutamakan {{prioritas_audiens}}.",
            "variables": ["nama_produk", "tema_konten", "target_audiens", "prioritas_audiens"],
            "example": {
                "input": {
                    "nama_produk": "Kemeja Flannel Tartan",
                    "nama_brand": "Nusa Studio",
                    "tema_konten": "Styling kampus vs nongkrong",
                    "target_audiens": "Mahasiswa dan pekerja muda",
                    "prioritas_audiens": "Kualitas kain tidak gatal dan potongan pas",
                },
                "output": "1. Jangan beli flannel sebelum tahu ciri kain yang tidak bikin gerah seharian.\n2. Satu kemeja, tiga gaya: rahasia outfit nongkrong langsung rapi.\n3. Alasan kami butuh 4 kali revisi pola demi potongan bahu yang tidak kaku ini.\n4. Ini upgrade terbaik untuk lemari pakaian kamu di bawah IDR 250.000.\n5. Stop pakai kemeja tipis yang mudah kusut sehabis dicuci sekali.",
            },
        },
        {
            "id": "brief-lookbook-photoshoot",
            "template": "Susun lembar brief pemotretan lookbook untuk koleksi {{nama_koleksi}} oleh {{nama_brand}}. Pemotretan bertempat di {{lokasi_shoot}} dengan suasana (mood) {{suasana_visual}}. Rinci konsep pencahayaan, arahan pose model, dan daftar outfit prioritas.",
            "variables": ["nama_koleksi", "nama_brand", "lokasi_shoot", "suasana_visual"],
            "example": {
                "input": {
                    "nama_koleksi": "Urban Transit Collection",
                    "nama_brand": "Vektor Works",
                    "lokasi_shoot": "Stasiun kereta dan trotoar perkotaan sore hari",
                    "suasana_visual": "Cahaya sore keemasan (golden hour), candid dinamis, dan tone warna hangat",
                },
                "output": "Brief Pemotretan Koleksi Urban Transit:\n1. Konsep Utama: Aktivitas harian komuter urban yang dinamis dan fungsional.\n2. Pencahayaan: Dominan natural golden hour sore hari dengan sedikit fill flash direct untuk efek kontras modern.\n3. Pose Model: Gerakan berjalan natural, menyeberang jalan, dan bersandar santai di railing tangga stasiun.\n4. Prioritas Foto: Detail tekstur kain, kerapian jahitan saku celana kargo, dan logo woven pada lengan baju.",
            },
        },
        {
            "id": "dm-seeding-influencer",
            "template": "Tuliskan draf pesan DM Instagram penawaran gifting/seeding produk {{nama_produk}} dari {{nama_brand}} kepada kreator {{nama_kreator}}. Jelaskan alasan brand memilih kreator tersebut dan ajakan kolaborasi tanpa paksaan berlebihan.",
            "variables": ["nama_produk", "nama_brand", "nama_kreator"],
            "example": {
                "input": {
                    "nama_produk": "Vintage Washed Hoodie",
                    "nama_brand": "Kroma Goods",
                    "nama_kreator": "Rian Kusuma",
                },
                "output": "Halo Rian! Kami dari tim Kroma Goods menyukai kurasi gaya pakaian harian kamu yang konsisten dan rapi. Dalam rangka rilis seri Vintage Washed Hoodie pekan ini, kami ingin mengirimkan satu paket lengkap untuk kamu coba langsung sesuai ukuranmu. Tidak ada kewajiban unggah khusus, kami hanya ingin kamu merasakan kenyamanan materialnya. Jika berkenan, boleh kirimkan alamat dan ukuran favoritmu ya. Terima kasih banyak!",
            },
        },
        {
            "id": "caption-teaser-drop",
            "template": "Tuliskan caption Instagram pengumuman rilis koleksi {{nama_koleksi}} dari {{nama_brand}}. Rilis dijadwalkan pada {{tanggal_rilis}} dengan jumlah terbatas sebanyak {{total_kuota}} pcs. Masukkan call-to-action untuk mengaktifkan notifikasi pengingat.",
            "variables": ["nama_koleksi", "nama_brand", "tanggal_rilis", "total_kuota"],
            "example": {
                "input": {
                    "nama_koleksi": "Monochrome Chapter II",
                    "nama_brand": "Sora Project",
                    "tanggal_rilis": "25 Oktober 2026 pukul 19.00 WIB",
                    "total_kuota": "150",
                },
                "output": "Eksplorasi siluet minimalis kembali berlanjut. Monochrome Chapter II resmi kami lepas pada 25 Oktober 2026 pukul 19.00 WIB eksklusif di situs resmi kami. Batch kali ini hanya diproduksi sebanyak 150 pcs tanpa rencana restock ulang. Ketuk stiker pengingat di postingan ini agar tidak kehabisan ukuran pilihanmu.",
            },
        },
        {
            "id": "cs-komplain-ukuran-salah",
            "template": "Buat template pesan balasan Customer Service WhatsApp saat pembeli mengeluh ukuran baju {{nama_produk}} yang diterima kekecilan. Tunjukkan rasa empati, jelaskan syarat penukaran ukuran, dan berikan panduan langkah pengiriman balik.",
            "variables": ["nama_produk"],
            "example": {
                "input": {"nama_produk": "Worker Overshirt"},
                "output": "Halo Kak, terima kasih sudah menghubungi kami. Mohon maaf sekali atas ketidaknyamanannya terkait ukuran Worker Overshirt yang kurang pas di badan Kakak. Kami menyediakan layanan tukar ukuran dalam 3 hari kerja sejak paket diterima dengan syarat label belum dilepas dan baju belum dicuci. Boleh informasikan ukuran pengganti yang diinginkan? Tim kami akan segera menyiapkan resi retur dan panduan pengiriman balik.",
            },
        },
    ]

    (folder / "prompts.json").write_text(json.dumps({"prompts": prompts}, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (folder / "README.md").write_text(
        "# Paket Prompt Pemasaran dan Peluncuran Clothing Brand\n\n"
        "Paket prompt terstruktur untuk brand pakaian, mencakup deskripsi katalog, naskah hook video, brief pemotretan, pendekatan influencer, dan penanganan layanan pelanggan.\n",
        encoding="utf-8",
    )
    return folder


def build_spreadsheet() -> Path:
    folder = BASE / "kalkulator-hpp-clothing"
    folder.mkdir(parents=True, exist_ok=True)

    metadata = {
        "title": "Kalkulator HPP dan Finansial Clothing Brand",
        "author": "Digital Product Factory",
        "version": "1.0.0",
        "license": "CC0-1.0",
        "input_cells": [
            "Kalkulator!C4", "Kalkulator!C5", "Kalkulator!C6", "Kalkulator!C7",
            "Kalkulator!C8", "Kalkulator!C9", "Kalkulator!C10", "Kalkulator!C11",
            "Kalkulator!C12", "Kalkulator!C16", "Kalkulator!C17",
            "Batch!C4", "Batch!C6", "Batch!C7", "Batch!C8", "Batch!C9"
        ],
    }
    (folder / "product.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    wb = openpyxl.Workbook()

    # Sheet 1: Instructions
    ws_inst = wb.active
    ws_inst.title = "Instructions"
    ws_inst.views.sheetView[0].showGridLines = True

    header_font = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
    bold_font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    normal_font = Font(name="Segoe UI", size=10, color="334155")
    title_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws_inst["A1"] = "PANDUAN KALKULATOR HPP & FINANSIAL CLOTHING BRAND"
    ws_inst["A1"].font = header_font

    instructions_data = [
        ("1. Tujuan Template", "Menghitung Harga Pokok Produksi (HPP) per helai baju secara presisi dan mensimulasikan titik impas (BEP) serta laba bersih per batch."),
        ("2. Lembar Kalkulator", "Masukkan biaya bahan baku, ongkos jahit, sablon/bordir, label, dan packaging pada sel berlatar warna kuning (Kolom C)."),
        ("3. Lembar Batch", "Tentukan volume produksi satu batch dan biaya peluncuran (sample, foto lookbook, iklan) untuk melihat proyeksi laba."),
        ("4. Aturan Input", "Isi hanya angka tanpa titik atau simbol mata uang pada sel input yang telah ditandai."),
        ("5. Formula Otomatis", "Kolom rekomendasi harga jual, margin, dan titik impas dihitung otomatis oleh formula spreadsheet."),
    ]

    for idx, (title, desc) in enumerate(instructions_data, start=3):
        ws_inst[f"A{idx}"] = title
        ws_inst[f"A{idx}"].font = bold_font
        ws_inst[f"B{idx}"] = desc
        ws_inst[f"B{idx}"].font = normal_font

    ws_inst.column_dimensions["A"].width = 28
    ws_inst.column_dimensions["B"].width = 85

    # Sheet 2: Kalkulator
    ws_calc = wb.create_sheet(title="Kalkulator")
    ws_calc.views.sheetView[0].showGridLines = True

    input_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
    calc_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    highlight_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    border_thin = Side(border_style="thin", color="CBD5E1")
    box_border = Border(top=border_thin, left=border_thin, right=border_thin, bottom=border_thin)

    ws_calc["A1"] = "KALKULATOR HPP SATUAN (PER HELAI KAOS/KEMEJA)"
    ws_calc["A1"].font = header_font

    calc_rows = [
        ("Komponen Biaya Produksi", "Keterangan", "Nilai (IDR)", "is_header"),
        ("Kain per Yard/Kg", "Harga beli per satuan bahan", 45000, "input"),
        ("Konsumsi Bahan per Helai", "Fraksi bahan (misal 0.35 kg/kaos)", 0.35, "input"),
        ("Ongkos Jahit Standar", "Ongkos potong dan jahit konveksi", 15000, "input"),
        ("Biaya Sablon / Bordir", "Biaya cetak grafik per titik", 12000, "input"),
        ("Label Leher (Main Label)", "Woven damask label leher", 1500, "input"),
        ("Label Samping (Washing)", "Satin washing instruction label", 800, "input"),
        ("Hangtag & Tali Pin", "Hangtag tebal 350 gsm + tali loop", 1200, "input"),
        ("Packaging Ziplock", "Plastik kemasan ziplock sablon", 2500, "input"),
        ("Polymailer Pengiriman", "Plastik kemasan amplop kurir luar", 800, "input"),
        ("Total HPP per Helai", "Total seluruh komponen modal langsung", "=ROUND((C4*C5)+SUM(C6:C12),0)", "calc"),
        ("", "", "", "blank"),
        ("Pengaturan Harga Jual", "Target Keuntungan & Strategi", "Nilai Target", "is_header"),
        ("Target Margin Kotor (%)", "Target rasio laba kotor yang diharapkan", 0.60, "input"),
        ("Estimasi Diskon Promosi (%)", "Penyangga diskon promo saat peluncuran", 0.10, "input"),
        ("Rekomendasi Harga Retail", "Harga jual normal sebelum potongan", "=ROUND(C13/(1-C16), -2)", "highlight"),
        ("Harga Promo Launching", "Harga diskon pada masa kampanye perdana", "=ROUND(C18*(1-C17), -2)", "calc"),
        ("Rekomendasi Harga Reseller", "Harga khusus pembelian grosir min 12 pcs", "=ROUND(C13*1.35, -2)", "calc"),
        ("Laba Kotor per Helai (Retail)", "Keuntungan kotor setiap produk terjual", "=C18-C13", "highlight"),
    ]

    for r_idx, (col_a, col_b, col_c, row_type) in enumerate(calc_rows, start=3):
        ca = ws_calc[f"A{r_idx}"]
        cb = ws_calc[f"B{r_idx}"]
        cc = ws_calc[f"C{r_idx}"]

        ca.value = col_a
        cb.value = col_b
        cc.value = col_c

        ca.font = normal_font
        cb.font = normal_font
        cc.font = normal_font
        ca.border = box_border
        cb.border = box_border
        cc.border = box_border

        if row_type == "is_header":
            ca.font = bold_font
            cb.font = bold_font
            cc.font = bold_font
            ca.fill = title_fill
            cb.fill = title_fill
            cc.fill = title_fill
        elif row_type == "input":
            cc.fill = input_fill
            cc.font = bold_font
            if isinstance(col_c, float) and col_c < 1.0:
                cc.number_format = "0.0%"
            else:
                cc.number_format = "#,##0"
        elif row_type == "calc":
            ca.font = bold_font
            cc.font = bold_font
            cc.fill = calc_fill
            cc.number_format = "#,##0"
        elif row_type == "highlight":
            ca.font = bold_font
            cc.font = bold_font
            cc.fill = highlight_fill
            cc.number_format = "#,##0"

    ws_calc.column_dimensions["A"].width = 30
    ws_calc.column_dimensions["B"].width = 46
    ws_calc.column_dimensions["C"].width = 20

    # Sheet 3: Batch
    ws_batch = wb.create_sheet(title="Batch")
    ws_batch.views.sheetView[0].showGridLines = True

    ws_batch["A1"] = "SIMULASI PRODUKSI PER BATCH & TITIK IMPAS (BEP)"
    ws_batch["A1"].font = header_font

    batch_rows = [
        ("Parameter Batch Produksi", "Keterangan", "Nilai", "is_header"),
        ("Jumlah Produksi per Batch", "Total kuantiti produksi (pcs)", 100, "input"),
        ("Total Modal Fisik Produk", "Dihitung dari HPP x Kuantiti batch", "=C4*Kalkulator!C13", "calc"),
        ("Biaya Pembuatan Pola & Sample", "Biaya riset pola pertama kali", 250000, "input"),
        ("Biaya Sesi Foto Lookbook", "Sewa studio & model katalog", 750000, "input"),
        ("Anggaran Iklan & Seeding", "Budget promosi digital dan influencer", 500000, "input"),
        ("Biaya Tak Terduga (Buffer)", "Cadangan operasional pengiriman & packing", 200000, "input"),
        ("Total Biaya Tetap (Fixed Cost)", "Jumlah seluruh pengeluaran di luar produksi fisik", "=SUM(C6:C9)", "calc"),
        ("Total Investasi 1 Batch", "Total modal siap pakai yang dibutuhkan", "=C5+C10", "highlight"),
        ("", "", "", "blank"),
        ("Analisis Titik Impas & Proyeksi Laba", "Perhitungan Finansial", "Hasil Analisis", "is_header"),
        ("Titik Impas / BEP (Unit)", "Jumlah baju yang harus terjual untuk balik modal", "=ROUNDUP(C10/(Kalkulator!C18-Kalkulator!C13),0)", "highlight"),
        ("Titik Impas / BEP (Rupiah)", "Omzet minimum agar tidak merugi", "=C14*Kalkulator!C18", "calc"),
        ("Target Omzet (Sold Out 100%)", "Total penerimaan jika seluruh batch habis terjual", "=C4*Kalkulator!C18", "calc"),
        ("Proyeksi Laba Bersih (Net)", "Laba bersih setelah dipotong seluruh modal", "=C16-C11", "highlight"),
        ("Estimasi ROI (%)", "Rasio pengembalian investasi batch", "=ROUND((C17/C11)*100,1)", "calc"),
    ]

    for r_idx, (col_a, col_b, col_c, row_type) in enumerate(batch_rows, start=3):
        ca = ws_batch[f"A{r_idx}"]
        cb = ws_batch[f"B{r_idx}"]
        cc = ws_batch[f"C{r_idx}"]

        ca.value = col_a
        cb.value = col_b
        cc.value = col_c

        ca.font = normal_font
        cb.font = normal_font
        cc.font = normal_font
        ca.border = box_border
        cb.border = box_border
        cc.border = box_border

        if row_type == "is_header":
            ca.font = bold_font
            cb.font = bold_font
            cc.font = bold_font
            ca.fill = title_fill
            cb.fill = title_fill
            cc.fill = title_fill
        elif row_type == "input":
            cc.fill = input_fill
            cc.font = bold_font
            cc.number_format = "#,##0"
        elif row_type == "calc":
            ca.font = bold_font
            cc.font = bold_font
            cc.fill = calc_fill
            cc.number_format = "#,##0"
        elif row_type == "highlight":
            ca.font = bold_font
            cc.font = bold_font
            cc.fill = highlight_fill
            cc.number_format = "#,##0"

    ws_batch.column_dimensions["A"].width = 34
    ws_batch.column_dimensions["B"].width = 48
    ws_batch.column_dimensions["C"].width = 20

    xlsx_path = folder / "kalkulator-clothing.xlsx"
    wb.save(xlsx_path)
    return folder


def build_ebook() -> Path:
    folder = BASE / "ebook-clothing-brand"
    folder.mkdir(parents=True, exist_ok=True)

    metadata = {
        "title": "Blueprint Membangun Clothing Brand Lokal",
        "author": "Digital Product Factory",
        "version": "1.0.0",
        "license": "CC0-1.0",
    }
    (folder / "product.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    md_content = """---
title: "Blueprint Membangun Clothing Brand Lokal"
author: "Digital Product Factory"
lang: "id"
date: "2026-10-20"
---

# Bab 1: Fondasi & Positioning Brand

Membangun clothing brand yang berumur panjang menuntut kejelasan identitas sejak hari pertama. Pasar pakaian jadi di Indonesia sangat dinamis, namun pemain yang bertahan bukan yang sekadar mencetak grafis pada kaos oblong, melainkan mereka yang mampu menawarkan nilai, narasi, dan konsistensi kualitas.

## Menentukan Niche Pasar

1. **Streetwear Kontemporer**: Menekankan grafis ekspresif, potongan oversized/boxy, dan detail berat kain (heavyweight).
2. **Minimalist Basic Essentials**: Fokus mutlak pada kenyamanan bahan katun premium, jahitan bersih, tanpa grafis mencolok, mengincar gaya hidup harian.
3. **Pakaian Fungsional & Komuter**: Memakai fitur saku ergonomis, ketahanan kain terhadap cuaca, dan fleksibilitas mobilitas perkotaan.

## Identitas Visual

Tentukan palet warna utama (maksimal 3 warna dasar), tipografi resmi brand, dan panduan foto katalog sebelum membuat materi promosi pertama Anda.

---

# Bab 2: Anatomi Bahan & Standar Kain Kaos

Pemilihan bahan adalah 70 persen penentu kepuasan pelanggan pakaian. Jangan hanya mengandalkan nama bahan dari perantara; pelajari langsung berat gramasi dan rajutannya.

## Klasifikasi Cotton Combed

- **Cotton Combed 30s (140-150 gsm)**: Karakter ringan, sejuk, sangat cocok untuk iklim kota panas dan pemakaian harian santai.
- **Cotton Combed 24s (175-185 gsm)**: Ketebalan menengah paling aman untuk pasar Indonesia. Tidak terlalu tipis dan tidak membuat gerah.
- **Cotton Combed 16s / Heavyweight (230-260 gsm)**: Struktur tebal dan kaku, membentuk siluet boxy fit yang modern dan tahan cuci berkali-kali.

## Karakter Penyusutan (Shrinkage)

Kain katun murni memiliki tingkat susut alami sekitar 2 hingga 4 persen setelah pencucian pertama. Selalu hitung toleransi susut ini pada pola mentah sebelum dipotong masal.

---

# Bab 3: Manajemen Vendor & Konveksi

Bekerja sama dengan konveksi atau penjahit memerlukan lembar kerja teknis (tech-pack) yang terperinci agar meminimalkan kesalahpahaman.

## Lembar Ukuran (Size Chart)

Tentukan standar ukuran brand Anda sendiri. Cantumkan toleransi jahitan garment plus minus 1 hingga 1.5 cm pada setiap komponen: lebar dada, panjang badan, dan panjang lengan.

## Metode Sablon Utama

1. **Plastisol**: Warna tajam, daya rekat tinggi pada serat kain, tahan lama, namun tidak dapat disetrika langsung pada permukaan tinta.
2. **Discharge (Cabut Warna)**: Menyerap ke dalam serat kain sehingga permukaan sablon terasa rata tanpa lapisan tebal (breathable).
3. **DTF (Direct to Film)**: Fleksibel untuk detail gradasi dan kuantiti kecil, namun rentan kaku jika area gambar terlalu lebar.

---

# Bab 4: Struktur HPP & Strategi Penetapan Harga

Keberlanjutan bisnis clothing bergantung pada perhitungan biaya yang mencakup komponen tak terlihat seperti kemasan dan penanganan susut.

## Komponen Modal Pokok (HPP)

1. Biaya Kain Pokok per Helai
2. Biaya Potong dan Jahit
3. Biaya Sablon / Bordir
4. Label Woven Leher dan Label Samping
5. Kemasan Ziplock dan Polymailer Pengiriman
6. Quality Control dan Pembersihan Benang

## Rumus Margin Sehat

- **Penjualan Langsung ke Konsumen (D2C)**: Target margin kotor ideal berkisar antara 55 hingga 65 persen. Margin ini menyerap biaya pemasaran digital dan promosi.
- **Penjualan Grosir / Reseller**: Tetapkan harga jual reseller dengan margin minimal 30 persen di atas HPP agar mitra memiliki ruang menjual kembali.

---

# Bab 5: Strategi Peluncuran Koleksi

Peluncuran koleksi pakaian membutuhkan penumpukan antisipasi audiens sebelum toko resmi dibuka.

## Rangkaian Kampanye

- **Minggu -2**: Unggah foto teaser potongan kain, siluet gelap, atau proses produksi di balik layar.
- **Minggu -1**: Publikasikan lookbook foto resmi yang memperlihatkan pakaian dikenakan model di lokasi tematik.
- **Hari H**: Buka penjualan tepat waktu dengan batas kuota stok yang jelas (metode drop) untuk menciptakan urgensi pembelian yang wajar.

---

# Bab 6: SOP Layanan Pelanggan & Retur

Pondasi reputasi brand fesyen pemula adalah cara menangani ketidaksesuaian ukuran pakaian yang diterima konsumen.

1. Siapkan batas waktu pelaporan penukaran maksimal 3 hari kerja sejak paket diterima kurir.
2. Syarat utama: label hangtag belum terlepas dan pakaian belum dicuci atau diberi pewangi tambahan.
3. Respon cepat dari tim layanan pelanggan mengubah komplain awal menjadi loyalitas pembeli setia.
"""

    md_path = folder / "book.md"
    md_path.write_text(md_content, encoding="utf-8")

    epub_path = folder / "book.epub"
    html_path = folder / "temp.html"
    pdf_path = folder / "book.pdf"

    # 1. Generate EPUB via pandoc
    subprocess.run(["pandoc", str(md_path), "-o", str(epub_path), "--toc", "--toc-depth=2"], check=True)

    # 2. Generate HTML with simple clean style for PDF
    css_content = """
    body { font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; line-height: 1.6; color: #1e293b; padding: 30px; }
    h1 { color: #0f172a; border-bottom: 2px solid #e2e8f0; padding-bottom: 8px; margin-top: 32px; }
    h2 { color: #334155; margin-top: 24px; }
    p, li { font-size: 11pt; color: #334155; }
    hr { border: none; border-top: 1px solid #cbd5e1; margin: 32px 0; }
    """
    html_wrapped = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Blueprint Membangun Clothing Brand Lokal</title><style>{css_content}</style></head><body>"
    html_raw = subprocess.run(["pandoc", str(md_path), "-t", "html", "--toc"], capture_output=True, text=True, check=True).stdout
    html_wrapped += html_raw + "</body></html>"
    html_path.write_text(html_wrapped, encoding="utf-8")

    # 3. Convert HTML to PDF via LibreOffice
    soffice = tool_path("soffice")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(folder), str(html_path)], check=True)
    temp_pdf = folder / "temp.pdf"
    if temp_pdf.is_file():
        temp_pdf.replace(pdf_path)

    md_path.unlink()
    html_path.unlink()
    return folder


def main() -> None:
    print("Membangun 3 produk digital unggulan...")
    pp_dir = build_prompt_pack()
    print("Prompt pack dibuat di:", pp_dir)
    rep_pp = check_product(pp_dir, "prompt-pack")
    print(f"Prompt pack QA exit: {rep_pp.exit_code} (findings: {len(rep_pp.findings)})")

    calc_dir = build_spreadsheet()
    print("Spreadsheet dibuat di:", calc_dir)
    rep_calc = check_product(calc_dir, "xlsx")
    print(f"Spreadsheet QA exit: {rep_calc.exit_code} (findings: {len(rep_calc.findings)})")

    ebook_dir = build_ebook()
    print("Ebook dibuat di:", ebook_dir)
    rep_ebook = check_product(ebook_dir, "ebook")
    print(f"Ebook QA exit: {rep_ebook.exit_code} (findings: {len(rep_ebook.findings)})")

    assert rep_pp.exit_code == 0, "Prompt pack QA failed"
    assert rep_calc.exit_code == 0, "Spreadsheet QA failed"
    assert rep_ebook.exit_code == 0, "Ebook QA failed"
    print("\nSEMUA 3 PRODUK LOLOS QA GATE DENGAN EXIT 0!")


if __name__ == "__main__":
    main()
